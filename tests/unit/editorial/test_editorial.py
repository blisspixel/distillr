"""Editorial budget, source, route, review and end-to-end failure contracts."""

from __future__ import annotations

import hashlib
import io
import json
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
from zipfile import ZipFile

import httpx
import pytest
from docx import Document
from pydantic import ValidationError
from typer.testing import CliRunner

from distill.editorial import budget, routing, sources, workflow
from distill.editorial.budget import EditorialTracker, WeeklyLedger, utc_week
from distill.editorial.perspective import (
    EditorialBudget,
    EditorialRoutes,
    EditorialSources,
    Perspective,
    init_perspective,
    load_perspective,
)
from distill.editorial.retonr import check_revision
from distill.editorial.review import Review, structural_issues, verify_review
from distill.ingestors.net import NetworkError
from distill.ingestors.podcasts.feed import PodcastEpisode, PodcastFeed
from distill.library.article_export import export_article
from distill.llm.cost_policy import CostPolicyError
from distill.llm.router import ConfigurationError, RouterConfig
from distill.llm.types import LLM_Response
from distill.llm.usage import LLMUsageAttempt
from distill.pipeline.budget import ProjectedBudgetExceededError

NOW = datetime(2026, 9, 9, 12, tzinfo=UTC)
URL = "https://example.org/news"
ARTICLE = f"# A practical AI update\n\nThe lab released a model. [Announcement]({URL})\n\n## What to try\n\nCompare it on a task you know."


@pytest.fixture
def source():
    return sources.receipt(
        url=URL,
        title="Model announcement",
        text="The lab released a model.",
        published=NOW.isoformat(),
        kind="feed",
    )


@pytest.fixture
def brief():
    return Perspective(topics=["AI news"])


def review_data(source, *, verdict="pass"):
    data = {
        key: {"verdict": verdict, "reason": "Supported by this source."}
        for key in (
            "grounding",
            "freshness",
            "perspective",
            "usefulness",
            "structure",
            "style",
            "coverage",
        )
    }
    data["claims"] = [
        {
            "claim": "The lab released a model.",
            "source_id": source.id,
            "quote": "The lab released a model.",
        }
    ]
    data["revision_instructions"] = "" if verdict == "pass" else "Clarify the takeaway."
    return data


def test_private_init_and_strict_config(tmp_path):
    path = init_perspective(tmp_path / "private" / "perspective.toml")
    example = load_perspective(path)
    assert example.author.company == ""
    assert not example.budget.allow_metered
    assert example.routes.refinement_model.startswith("openai/")
    with pytest.raises(FileExistsError):
        init_perspective(path)
    path.write_text("x" * 64001)
    with pytest.raises(ValueError, match="64000"):
        load_perspective(path)


@pytest.mark.parametrize(
    "changes",
    [
        {"weekly_usd": -1.0},
        {"weekly_usd": float("nan")},
        {"weekly_usd": float("inf")},
        {"per_run_usd": True},
        {"per_run_usd": 6.0},
    ],
)
def test_bad_budgets(changes):
    with pytest.raises(ValueError):
        EditorialBudget(**changes)


@pytest.mark.parametrize(
    "url",
    [
        "http://example.org/a",
        "https://u:p@example.org",
        "https://example.org/#x",
        "https://example.org:0/",
        "https://example.org/a b",
        "https:///bad",
        "https://example.org/" + "x" * 2100,
    ],
)
def test_bad_source_urls(url):
    with pytest.raises(ValueError):
        EditorialSources(urls=[url])


def test_strict_models_and_routes():
    with pytest.raises(ValidationError):
        Perspective(topics=["AI"], unknown=True)
    with pytest.raises(ValueError):
        EditorialRoutes(refinement_model="anthropic/claude-sonnet-5")
    with pytest.raises(ValueError):
        EditorialRoutes(writer_model="openrouter/auto")
    assert EditorialSources(urls=[URL, URL]).urls == [URL]
    assert utc_week(NOW) == "2026-09-07"
    assert utc_week(NOW - timedelta(days=3)) == "2026-08-31"


def test_weekly_ledger_restart_rollover_and_corruption(tmp_path, monkeypatch):
    path = tmp_path / "budget.jsonl"
    ledger = WeeklyLedger(path)
    monkeypatch.setattr(budget, "utc_week", lambda: "2026-09-07")
    first = ledger.reserve(4.0, 5.0, "one")
    with pytest.raises(ProjectedBudgetExceededError):
        WeeklyLedger(path).reserve(2.0, 5.0, "two")
    ledger.settle(first, 1.0)
    pending = ledger.reserve(3.0, 5.0, "two")
    assert ledger.status(5.0)["remaining_usd"] == 1.0
    monkeypatch.setattr(budget, "utc_week", lambda: "2026-09-14")
    assert ledger.status(5.0)["spent_or_reserved_usd"] == 3.0
    ledger.settle(pending, 0.5)
    assert ledger.status(5.0)["spent_or_reserved_usd"] == 0.5
    with pytest.raises(ValueError, match="already settled"):
        ledger.settle(pending, 0.5)
    with path.open("a") as stream:
        stream.write('{"broken":')
    with pytest.raises(ValueError):
        ledger.reserve(0.01, 5.0, "three")


def test_concurrent_reservations_cannot_overspend(tmp_path):
    path = tmp_path / "budget.jsonl"

    def reserve(_):
        try:
            WeeklyLedger(path).reserve(0.75, 1.0, "run")
            return True
        except ProjectedBudgetExceededError:
            return False

    with ThreadPoolExecutor(max_workers=4) as pool:
        assert sum(pool.map(reserve, range(4))) == 1


@pytest.mark.parametrize("damage", ["duplicate", "unmatched", "wrong_run"])
def test_invalid_ledger_transitions_fail_closed(tmp_path, damage):
    ledger = WeeklyLedger(tmp_path / "budget.jsonl")
    event = ledger.reserve(0.1, 1.0, "run")
    if damage == "duplicate":
        extra = event
    elif damage == "unmatched":
        extra = event.model_copy(update={"kind": "settle", "reservation": "f" * 32})
    else:
        extra = event.model_copy(update={"kind": "settle", "run_id": "other"})
    with ledger.path.open("a") as stream:
        stream.write(extra.model_dump_json() + "\n")
    with pytest.raises(ValueError):
        ledger.status(1.0)


@pytest.mark.parametrize("outcome", ["success", "error", "crash", "zero"])
def test_tracker_reserves_before_contact_and_accounts_failures(tmp_path, outcome):
    ledger = WeeklyLedger(tmp_path / "budget.jsonl")
    tracker = EditorialTracker(ledger, weekly_usd=1.0, per_run_usd=0.5, run_id="run")
    attempt = LLMUsageAttempt(
        input_tokens=1000,
        output_tokens=1000,
        model="anthropic/claude-sonnet-5",
        provider_name="openrouter",
        provider_type="cloud",
        usage_source="reported",
        outcome="success",
    )

    def invoke():
        with tracker.reserve_attempt(attempt):
            assert ledger.status(1.0)["pending_reservations"] == 1
            if outcome == "crash":
                raise KeyboardInterrupt()
            if outcome != "zero":
                tracker.record_attempt(attempt)
            if outcome == "error":
                raise RuntimeError("provider error")

    if outcome in {"error", "crash"}:
        with pytest.raises(RuntimeError if outcome == "error" else KeyboardInterrupt):
            invoke()
    else:
        invoke()
    assert ledger.status(1.0)["pending_reservations"] == (1 if outcome == "crash" else 0)
    assert tracker.total_cost <= 0.5


def test_tracker_run_limit_refuses_before_weekly_reservation(tmp_path):
    ledger = WeeklyLedger(tmp_path / "budget.jsonl")
    tracker = EditorialTracker(ledger, weekly_usd=1, per_run_usd=0, run_id="run")
    attempt = LLMUsageAttempt(
        input_tokens=1000,
        output_tokens=1000,
        model="anthropic/claude-sonnet-5",
        provider_name="openrouter",
        provider_type="cloud",
        usage_source="reported",
        outcome="success",
    )
    with pytest.raises(ProjectedBudgetExceededError), tracker.reserve_attempt(attempt):
        pytest.fail("provider contact")
    assert not ledger.path.exists()


def test_local_preference_cannot_fall_through_to_spend(monkeypatch, brief):
    base = RouterConfig(_env_file=None, cost_mode="no-metered")
    monkeypatch.setattr(routing, "local_available", lambda *a: True)
    brief.routes.local_model = "installed-model"
    configs, _ = routing.select_routes(brief, base)
    assert configs["writer"].resolve("report") == ("ollama", "installed-model")
    assert configs["writer"].fallback_provider == ""
    monkeypatch.setattr(routing, "local_available", lambda *a: False)
    with pytest.raises(CostPolicyError):
        routing.select_routes(brief, base)
    brief.budget.allow_metered = True
    with pytest.raises(CostPolicyError):
        routing.select_routes(brief, base, paid_only=True)
    configs, reason = routing.select_routes(
        brief, base.model_copy(update={"cost_mode": "paid-ok", "openrouter_api_key": "test"})
    )
    assert configs["refiner"].provider == "openrouter"
    assert configs["refiner"].openrouter_zdr is True
    assert "weekly" in reason
    monkeypatch.setattr(routing, "has_known_pricing", lambda _: False)
    with pytest.raises(CostPolicyError, match="registered"):
        routing.select_routes(brief, base.model_copy(update={"cost_mode": "paid-ok"}))


@pytest.mark.parametrize("provider,key", [("ollama", "models"), ("lmstudio", "data")])
def test_local_probe(monkeypatch, provider, key):
    row_key = "name" if provider == "ollama" else "id"
    response = Mock()
    response.json.return_value = {key: [{row_key: "installed"}]}
    get = Mock(return_value=response)
    monkeypatch.setattr(httpx, "get", get)
    assert routing.local_available(provider, "installed")
    assert get.call_args.kwargs["follow_redirects"] is False
    assert get.call_args.kwargs["trust_env"] is False
    assert not routing.local_available(provider, "missing")
    assert not routing.local_available(provider, "")
    get.side_effect = httpx.ConnectError("offline")
    assert not routing.local_available(provider, "installed")
    monkeypatch.setattr(routing, "classify_provider", lambda _: "unknown")
    assert not routing.local_available(provider, "installed")


def episode(url, published, text="A source claim."):
    return PodcastEpisode(
        title="News",
        guid=url,
        published=published,
        audio_url="",
        audio_type="",
        duration_s=0,
        description=text,
        link=url,
    )


def test_current_sources_keep_dates_dedupe_and_capture_failures(tmp_path, monkeypatch, source):
    items = [
        episode(URL, "Wed, 09 Sep 2026 01:00:00 GMT"),
        episode(URL, "Wed, 09 Sep 2026 01:00:00 GMT"),
        episode(URL + "/old", "Wed, 01 Jan 2020 01:00:00 GMT"),
        episode(URL + "/future", "Thu, 10 Sep 2026 01:00:00 GMT"),
        episode(URL + "/undated", ""),
        episode(URL + "/empty", "Wed, 09 Sep 2026 01:00:00 GMT", ""),
    ]

    def feed(url):
        if "broken" in url:
            raise NetworkError("unavailable")
        return PodcastFeed(title="Publisher", link=URL, description="", episodes=items)

    monkeypatch.setattr(sources, "fetch_feed", feed)
    monkeypatch.setattr(sources, "fetch_page", Mock(side_effect=NetworkError("offline")))
    config = EditorialSources(feeds=[URL, URL + "/broken"], urls=[URL + "/page"])
    captured, failures = sources.collect_sources(config, now=NOW)
    assert len(captured) == 1
    assert len(failures) == 3
    assert captured[0].published_at.startswith("2026-09-09")
    sources.save_receipts(tmp_path, captured)
    assert json.loads(next(tmp_path.glob("*.json")).read_text())["text"] == "A source claim."
    assert sources.evidence_packet(captured, 1)[0]["excerpt_truncated"] is True
    monkeypatch.setattr(sources, "fetch_page", lambda _: source)
    captured, _ = sources.collect_sources(
        EditorialSources(feeds=[URL], urls=[URL], max_sources=2), now=NOW
    )
    assert len(captured) == 1
    captured, _ = sources.collect_sources(EditorialSources(feeds=[URL], max_sources=1), now=NOW)
    assert captured[0].text == source.text


def test_capture_page_size_and_redirect(monkeypatch):
    class Response(io.BytesIO):
        def geturl(self):
            return URL

    monkeypatch.setattr(sources, "safe_urlopen", lambda *a, **k: Response(b"<p>Source text.</p>"))
    captured = sources.fetch_page(URL)
    assert captured.url == URL and "Source text." in captured.text
    monkeypatch.setattr(sources, "safe_urlopen", lambda *a, **k: Response(b"x" * 4000001))
    with pytest.raises(ValueError, match="4 MB"):
        sources.fetch_page(URL)
    assert sources.receipt(url=URL, title="Title", text="x" * 120001, kind="page").truncated


@pytest.mark.parametrize(
    "damage",
    [
        "none",
        "empty",
        "oversize",
        "digest",
        "identity",
        "duplicate",
        "stale",
        "naive",
        "future",
        "url",
    ],
)
def test_supplied_receipt_integrity_and_freshness(tmp_path, source, damage):
    source = source.model_copy(update={"fetched_at": datetime.now(UTC).isoformat()})
    sources.save_receipts(tmp_path, [source])
    path = tmp_path / f"{source.id}.json"
    data = source.model_dump()
    if damage == "empty":
        path.unlink()
    elif damage == "oversize":
        path.write_text("x" * 600001)
    elif damage == "duplicate":
        (tmp_path / "duplicate.json").write_text(path.read_text())
    elif damage != "none":
        field, value = {
            "digest": ("sha256", "0" * 64),
            "identity": ("id", "s" + "0" * 16),
            "stale": ("fetched_at", "2000-01-01T00:00:00+00:00"),
            "naive": ("fetched_at", "2026-09-09T00:00:00"),
            "future": ("fetched_at", "2999-01-01T00:00:00+00:00"),
            "url": ("url", "http://example.org/news"),
        }[damage]
        data[field] = value
        path.write_text(json.dumps(data), encoding="utf-8")
    if damage == "none":
        result = sources.load_receipts(tmp_path, EditorialSources())
        assert result[0].capture_method == "supplied-receipt"
        assert result[0].sha256 == source.sha256
    else:
        with pytest.raises(ValueError):
            sources.load_receipts(tmp_path, EditorialSources())


def test_supplied_receipt_workflow_and_truncated_response(tmp_path, source, brief, monkeypatch):
    fake_workflow(monkeypatch, source)
    receipt_dir = tmp_path / "input"
    sources.save_receipts(receipt_dir, [source])
    folder = workflow.build_article(brief, tmp_path, RouterConfig(), receipt_dir=receipt_dir)
    assert (folder / "article.md").exists()
    monkeypatch.setattr(
        workflow,
        "call",
        lambda *a, **k: LLM_Response(
            text="Partial output",
            input_tokens=1,
            output_tokens=8192,
            model="test",
            finish_reason="length",
        ),
    )
    with pytest.raises(workflow.EditorialReviewError, match="Incomplete"):
        workflow.build_article(brief, tmp_path, RouterConfig(), receipt_dir=receipt_dir)


def test_capture_cli_and_ready_preview(tmp_path, source, monkeypatch):
    from distill.cli import app
    from distill.commands import editorial

    monkeypatch.chdir(tmp_path)
    init_perspective(Path("private/perspective.toml"))
    monkeypatch.setattr(editorial, "collect_sources", lambda _: ([source], []))
    route = RouterConfig(provider="ollama", model="installed")
    monkeypatch.setattr(editorial, "select_routes", lambda *a: ({"writer": route}, "Local"))
    runner = CliRunner()
    result = runner.invoke(app, ["--json", "editorial", "preview"])
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)["data"]["ready"] is True
    result = runner.invoke(app, ["--json", "editorial", "capture", "private/packet"])
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)["data"]["source_count"] == 1
    assert runner.invoke(app, ["editorial", "capture", "private/packet"]).exit_code != 0


def test_review_exact_receipts_and_model_verdict(source):
    review = Review.model_validate(review_data(source))
    assert review.accepted and verify_review(review, ARTICLE, [source]) == []
    assert not Review.model_validate(review_data(source, verdict="revise")).accepted
    review.claims[0].quote = "fabricated quote"
    review.claims[0].claim = "not in article"
    assert len(verify_review(review, ARTICLE, [source])) == 2
    review.claims[0].quote = source.text
    assert verify_review(review, ARTICLE.replace(URL, URL + "/invented"), [source])


@pytest.mark.parametrize(
    "article",
    [
        "No title",
        ARTICLE + "\n# Another title",
        ARTICLE + "\u2014",
        ARTICLE + "\U0001f600",
        ARTICLE.replace(URL, "https://evil.example"),
        "# Title\n\nNo citations",
        ARTICLE + "\n```code```",
        ARTICLE + "\n<script>bad</script>",
        ARTICLE + "\n\n## A heading, with punctuation\n",
    ],
)
def test_output_shape_violations(article, source):
    assert structural_issues(article, [source])


def fake_workflow(monkeypatch, source, *, reviews=None, error=None):
    route = RouterConfig(_env_file=None, provider="ollama", model="test")
    monkeypatch.setattr(
        workflow,
        "select_routes",
        lambda *a, **k: (
            dict.fromkeys(["research", "writer", "refiner", "critic"], route),
            "Local",
        ),
    )
    monkeypatch.setattr(workflow, "collect_sources", lambda _: ([source], []))
    calls = []

    def model_call(config, workload, prompt, **kwargs):
        stage = kwargs["call_type"].removeprefix("editorial_")
        calls.append(stage)
        if error:
            raise error
        if stage in {"review", "repair_review"}:
            text = json.dumps(reviews.pop(0) if reviews else review_data(source))
        elif stage in {"draft", "refine", "revise"}:
            text = ARTICLE
        else:
            text = "Useful notes with a concrete argument."
        return LLM_Response(text=text, input_tokens=1, output_tokens=1, model="test")

    monkeypatch.setattr(workflow, "call", model_call)
    return calls


def test_end_to_end_requires_rewrite_review_and_preserves_history(
    tmp_path, monkeypatch, source, brief
):
    calls = fake_workflow(
        monkeypatch, source, reviews=[review_data(source, verdict="revise"), review_data(source)]
    )
    progress = []
    folder = workflow.build_article(brief, tmp_path, RouterConfig(), progress=progress.append)
    assert calls == ["reading", "ideas", "outline", "draft", "refine", "review", "revise", "review"]
    assert (folder / "article.md").read_text().strip() == ARTICLE
    assert (folder / "article.docx").exists()
    manifest = json.loads((folder / "manifest.json").read_text())
    assert manifest["status"] == "ready_for_review" and manifest["published"] is False
    assert workflow._history(tmp_path) == ["A practical AI update"]
    assert (folder / "receipts" / f"{source.id}.json").exists()
    assert (folder / "draft.md").exists() and (folder / "rewrite.md").exists()


def test_review_copy_errors_get_one_bounded_critic_repair(tmp_path, monkeypatch, source, brief):
    invalid = review_data(source)
    invalid["claims"][0]["quote"] = "An invented or miscopied quotation"
    calls = fake_workflow(monkeypatch, source, reviews=[invalid, review_data(source)])
    folder = workflow.build_article(brief, tmp_path, RouterConfig())
    assert calls[-2:] == ["review", "repair_review"]
    assert "revise" not in calls
    assert (folder / "review-0-repair.json").exists()
    assert json.loads((folder / "checks-0.json").read_text())["accepted"]


@pytest.mark.parametrize(
    "failure", ["review", "no_sources", "provider", "empty", "bad_json", "retonr"]
)
def test_failed_work_is_never_presented_as_a_final_article(
    tmp_path, monkeypatch, source, brief, failure
):
    fake_workflow(
        monkeypatch,
        source,
        reviews=[review_data(source, verdict="abstain")] * 3,
        error=RuntimeError("secret provider data") if failure == "provider" else None,
    )
    if failure == "no_sources":
        monkeypatch.setattr(workflow, "collect_sources", lambda _: ([], []))
    if failure in {"empty", "bad_json"}:
        monkeypatch.setattr(
            workflow,
            "call",
            lambda *a, **k: LLM_Response(
                text="" if failure == "empty" else "malformed",
                input_tokens=0,
                output_tokens=0,
                model="test",
            ),
        )
    if failure == "retonr":
        fake_workflow(monkeypatch, source)
        monkeypatch.setattr(workflow, "check_revision", Mock(side_effect=ValueError("abstained")))
    with pytest.raises((ValueError, RuntimeError)):
        workflow.build_article(
            brief,
            tmp_path,
            RouterConfig(),
            retonr_executable=Path("retonr") if failure == "retonr" else None,
        )
    folder = next((tmp_path / "runs").iterdir())
    assert not (folder / "article.md").exists()
    raw = (folder / "manifest.json").read_text()
    assert json.loads(raw)["status"] == "failed"
    assert "secret provider data" not in raw


def test_docx_single_title_links_and_no_branding(tmp_path):
    path = tmp_path / "article.md"
    path.write_text(
        ARTICLE + "\n\n- **Bold** and *italic*\n\n1. First\n\n### Detail\n\nLine  \nbreak\n",
        encoding="utf-8",
    )
    docx = export_article(path)
    doc = Document(docx)
    assert sum(p.style.name == "Title" for p in doc.paragraphs) == 1
    assert doc.core_properties.author == ""
    assert not any("Distill" in p.text for p in doc.paragraphs)
    with ZipFile(docx) as zipped:
        assert URL in zipped.read("word/_rels/document.xml.rels").decode()
        assert "w:hyperlink" in zipped.read("word/document.xml").decode()
        assert "w:pBdr" not in zipped.read("word/styles.xml").decode()
    with pytest.raises(FileExistsError):
        export_article(path)


@pytest.mark.parametrize(
    "mode", ["pass", "abstain", "bad_schema", "not_object", "script", "bad_record", "bad_digest"]
)
def test_retonr_adapter_fails_closed_and_does_not_inherit_keys(tmp_path, monkeypatch, mode):
    from distill.editorial import retonr

    binary = tmp_path / ("retonr.cmd" if mode == "script" else "retonr.exe")
    binary.touch()
    payload = {
        "schema_version": 2 if mode == "bad_schema" else 1,
        "command": "check",
        "status": "ok",
        "result": {
            "schema_version": 2,
            "status": "failed" if mode == "bad_record" else "rewritten",
            "source_digest": "bad"
            if mode == "bad_digest"
            else hashlib.sha256(b"draft").hexdigest(),
            "output_digest": hashlib.sha256(b"candidate").hexdigest(),
        },
    }
    run = Mock(
        return_value=SimpleNamespace(
            returncode=3 if mode == "abstain" else 0,
            stdout=json.dumps([] if mode == "not_object" else payload).encode(),
        )
    )
    monkeypatch.setattr(retonr.subprocess, "run", run)
    monkeypatch.setenv("OPENROUTER_API_KEY", "secret")
    args = (binary, tmp_path / "draft.md", tmp_path / "candidate.md", tmp_path / "check.json")
    args[1].write_bytes(b"draft")
    args[2].write_bytes(b"candidate")
    if mode == "pass":
        check_revision(*args)
        assert "OPENROUTER_API_KEY" not in run.call_args.kwargs["env"]
        assert "--fail-on-abstain" in run.call_args.args[0]
    else:
        with pytest.raises(ValueError):
            check_revision(*args)


def test_editorial_cli_init_preview_and_run(tmp_path, monkeypatch, source):
    from distill.cli import app

    monkeypatch.chdir(tmp_path)
    runner = CliRunner()
    result = runner.invoke(app, ["--json", "editorial", "init"])
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)["data"]["metered_enabled"] is False
    monkeypatch.setattr(routing, "local_available", lambda *a: False)
    result = runner.invoke(app, ["--json", "editorial", "preview"])
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)["data"]["ready"] is False
    fake_workflow(monkeypatch, source)
    result = runner.invoke(app, ["--json", "editorial", "run"])
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)["data"]["status"] == "ready_for_review"


@pytest.mark.parametrize("json_mode", [True, False])
def test_cli_errors_and_budget_exit(tmp_path, monkeypatch, json_mode):
    from distill.cli import app
    from distill.commands import editorial

    monkeypatch.chdir(tmp_path)
    runner = CliRunner()
    prefix = ["--json"] if json_mode else []
    assert runner.invoke(app, [*prefix, "editorial", "init"]).exit_code == 0
    assert runner.invoke(app, [*prefix, "editorial", "init"]).exit_code != 0
    monkeypatch.setattr(
        editorial, "build_article", Mock(side_effect=ProjectedBudgetExceededError(1, 0))
    )
    result = runner.invoke(app, [*prefix, "editorial", "run"])
    assert result.exit_code == 6
    from distill.llm.providers.openrouter import OpenRouterRequestError

    monkeypatch.setattr(editorial, "build_article", Mock(side_effect=OpenRouterRequestError(404)))
    result = runner.invoke(app, [*prefix, "editorial", "run"])
    assert result.exit_code == 3
    assert "HTTP 404" in result.output
    monkeypatch.setattr(editorial, "load_perspective", Mock(side_effect=ConfigurationError("key")))
    assert runner.invoke(app, [*prefix, "editorial", "preview"]).exit_code == 3
