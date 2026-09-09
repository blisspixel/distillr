"""Unbranded article DOCX with real hyperlinks and a single title."""

from __future__ import annotations

from html.parser import HTMLParser
from pathlib import Path
from typing import cast

import markdown
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.styles.style import ParagraphStyle


def _styles(doc: DocxDocument) -> None:
    # Some Word templates include decorative rules in Title or Subtitle.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name, size in (("Normal", 11), ("Title", 24), ("Heading 1", 15), ("Heading 2", 13)):
        style = cast(ParagraphStyle, doc.styles[name])
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.12
        style.paragraph_format.widow_control = True
        if name != "Normal":
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(14)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""


class _ArticleParser(HTMLParser):
    def __init__(self, doc: DocxDocument):
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self.paragraph = None
        self.bold = False
        self.italic = False
        self.href = ""
        self.ordered = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        styles = {"h1": "Title", "h2": "Heading 1", "h3": "Heading 2", "p": "Normal"}
        if tag in styles:
            self.paragraph = self.doc.add_paragraph(style=styles[tag])
        elif tag == "li":
            self.paragraph = self.doc.add_paragraph(
                style="List Number" if self.ordered else "List Bullet"
            )
        elif tag in {"ol", "ul"}:
            self.ordered = tag == "ol"
        elif tag in {"strong", "em"}:
            setattr(self, "bold" if tag == "strong" else "italic", True)
        elif tag == "a":
            self.href = dict(attrs).get("href") or ""
        elif tag == "br" and self.paragraph is not None:
            self.paragraph.add_run().add_break()

    def handle_endtag(self, tag: str) -> None:
        if tag in {"strong", "em"}:
            setattr(self, "bold" if tag == "strong" else "italic", False)
        elif tag == "a":
            self.href = ""
        elif tag in {"h1", "h2", "h3", "p", "li"}:
            self.paragraph = None

    def handle_data(self, data: str) -> None:
        if self.paragraph is None:
            return
        if self.href.startswith("https://"):
            link = OxmlElement("w:hyperlink")
            link.set(
                qn("r:id"), self.paragraph.part.relate_to(self.href, RT.HYPERLINK, is_external=True)
            )
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "205B82")
            properties.append(color)
            run.append(properties)
            text = OxmlElement("w:t")
            text.set(qn("xml:space"), "preserve")
            text.text = data
            run.append(text)
            link.append(run)
            self.paragraph._p.append(link)
        else:
            run = self.paragraph.add_run(data)
            run.bold, run.italic = self.bold, self.italic


def export_article(md_path: Path, docx_path: Path | None = None) -> Path:
    """Export the editorial Markdown subset without report branding or a cover page."""
    destination = docx_path or md_path.with_suffix(".docx")
    if destination.exists():
        raise FileExistsError(destination)
    doc = Document()
    _styles(doc)
    text = md_path.read_text(encoding="utf-8")
    _ArticleParser(doc).feed(markdown.markdown(text))
    doc.core_properties.title = text.splitlines()[0].removeprefix("# ")
    doc.save(str(destination))
    return destination
